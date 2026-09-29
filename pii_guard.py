#!/usr/bin/env python3
"""AI에 넘기기 전 개인정보 마스킹 (Word / Excel / PDF / 텍스트) + 선택적 비밀번호 설정.

설치:  pip install python-docx openpyxl pypdf cryptography msoffcrypto-tool
사용:  python pii_guard.py 파일1 [파일2 ...] [--words 홍길동,김영희] [--star] [--map] [--lock | --lock-only]

  --words    자동으로 못 찾는 이름·회사명 등 (쉼표 구분)
  --star     토큰 대신 별표(*)로 가림 (복원 불가)
  --map      토큰→원본 대응표(_map.json) 저장. 이 파일에는 개인정보가 들어 있으니 따로 보관하세요.
  --lock       원본 파일(개인정보 포함)에 열기 비밀번호를 걸어 '이름_locked.확장자'로 함께 저장 (docx, xlsx, pdf)
  --lock-only  마스킹은 하지 않고 비밀번호만 설정
결과: 원본은 건드리지 않고 '이름_masked.확장자' 로 저장합니다. PDF는 '이름_pdf_masked.txt'(텍스트만 추출).
"""
import argparse, getpass, json, os, re, time

N, X = r'(?<!\d)', r'(?!\d)'
RULES = [
    ('주민등록번호', re.compile(N + r'\d{6}[- ]?[1-4]\d{6}' + X)),
    ('카드번호',     re.compile(N + r'\d{4}[- ]?\d{4}[- ]?\d{4}[- ]?\d{4}' + X)),
    ('휴대전화',     re.compile(N + r'01[016789][-. ]?\d{3,4}[-. ]?\d{4}' + X)),
    ('전화번호',     re.compile(N + r'0(?:2|[3-6]\d)[-. ]?\d{3,4}[-. ]?\d{4}' + X)),
    ('이메일',       re.compile(r'[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+')),
    ('사업자번호',   re.compile(N + r'\d{3}-\d{2}-\d{5}' + X)),
    ('여권번호',     re.compile(r'(?<![A-Za-z0-9])[MSROD]\d{8}(?!\d)')),
    ('IP주소',       re.compile(r'(?<![\d.])(?:\d{1,3}\.){3}\d{1,3}(?![\d.])')),
    ('계좌번호',     re.compile(N + r'\d{2,6}-\d{2,6}-\d{2,8}(?:-\d{1,4})?' + X)),
]
DATE = re.compile(r'(19|20)\d{2}-\d{2}-\d{2}')  # 날짜를 계좌번호로 오인하지 않도록


class Masker:
    def __init__(self, words=(), star=False):
        self.star, self.map, self.seen, self.cnt, self.found = star, {}, {}, {}, {}
        self.rules = list(RULES)
        ws = sorted({w.strip() for w in words if w.strip()}, key=len, reverse=True)
        if ws:
            self.rules.insert(0, ('지정어', re.compile('|'.join(map(re.escape, ws)))))

    def text(self, s):
        for label, rx in self.rules:
            s = rx.sub(lambda m, l=label: self._sub(l, m.group(0)), s)
        return s

    def _sub(self, label, m):
        if label == '계좌번호' and DATE.fullmatch(m):
            return m
        self.found[label] = self.found.get(label, 0) + 1
        if self.star:
            return re.sub(r'[0-9A-Za-z가-힣]', '*', m)
        key = (label, m)
        if key not in self.seen:
            n = self.cnt[label] = self.cnt.get(label, 0) + 1
            self.seen[key] = f'[{label}_{n}]'
            self.map[self.seen[key]] = m
        return self.seen[key]


def do_docx(src, dst, M):
    import docx
    d = docx.Document(src)

    def para(p):
        t = p.text
        m = M.text(t)
        if m != t and p.runs:          # 번호가 여러 run에 쪼개져 있어도 잡도록 문단 단위로 처리
            p.runs[0].text = m
            for r in p.runs[1:]:
                r.text = ''

    def block(c):
        for p in c.paragraphs:
            para(p)
        for tb in c.tables:
            for row in tb.rows:
                for cell in row.cells:
                    block(cell)

    block(d)
    for s in d.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer,
                     s.even_page_header, s.even_page_footer):
            block(part)
    d.save(dst)


def do_xlsx(src, dst, M):
    import openpyxl
    wb = openpyxl.load_workbook(src)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and not v.startswith('='):
                    c.value = M.text(v)
                elif isinstance(v, int) and not isinstance(v, bool):
                    m = M.text(str(v))   # 숫자로 저장된 주민번호·카드번호 등
                    if m != str(v):
                        c.value = m
    wb.save(dst)


def do_pdf(src, dst, M):
    from pypdf import PdfReader
    text = '\n\n'.join((p.extract_text() or '') for p in PdfReader(src).pages)
    if not text.strip():
        raise ValueError('텍스트를 읽을 수 없는 PDF입니다(스캔본이면 OCR이 필요합니다).')
    # 줄바꿈으로 끊긴 이메일·번호를 이어 붙여서 놓치지 않게 함
    text = re.sub(r'(?<=[-@])\n(?=[A-Za-z0-9])|(?<=\.)\n(?=[a-z])|(?<=[A-Za-z0-9])\n(?=@)', '', text)
    open(dst, 'w', encoding='utf-8').write(M.text(text))


def do_text(src, dst, M):
    open(dst, 'w', encoding='utf-8').write(M.text(open(src, encoding='utf-8-sig').read()))


def lock(path, pw):
    """파일 자체에 열기 비밀번호를 설정 (Word/Excel/PDF 프로그램에서 바로 열림)."""
    base, ext = os.path.splitext(path)
    ext = ext.lower()
    dst = base + '_locked' + ext
    if ext == '.pdf':
        from pypdf import PdfReader, PdfWriter
        w = PdfWriter(clone_from=PdfReader(path))
        w.encrypt(pw, algorithm='AES-256')
        with open(dst, 'wb') as o:
            w.write(o)
    elif ext in ('.docx', '.xlsx'):
        from msoffcrypto.format.ooxml import OOXMLFile
        with open(path, 'rb') as i, open(dst, 'wb') as o:
            OOXMLFile(i).encrypt(pw, o)
    else:
        raise ValueError('비밀번호 설정은 docx, xlsx, pdf만 지원합니다.')
    return dst


def ask_pw():
    pw = getpass.getpass('파일 열기 비밀번호(8자 이상): ')
    if len(pw) < 8:
        raise SystemExit('비밀번호는 8자 이상이어야 합니다.')
    if pw != getpass.getpass('비밀번호 확인: '):
        raise SystemExit('비밀번호가 서로 다릅니다.')
    return pw


def main():
    ap = argparse.ArgumentParser(description='개인정보 마스킹 + 파일 비밀번호 설정')
    ap.add_argument('files', nargs='+')
    ap.add_argument('--words', default='')
    ap.add_argument('--star', action='store_true')
    ap.add_argument('--map', action='store_true')
    ap.add_argument('--lock', action='store_true')
    ap.add_argument('--lock-only', action='store_true')
    a = ap.parse_args()

    pw = ask_pw() if (a.lock or a.lock_only) else None

    for f in a.files:
        if a.lock_only:
            try:
                print(f'{f} -> {lock(f, pw)}')
            except Exception as e:
                print(f'실패: {f} - {e}')
            continue
        base, ext = os.path.splitext(f)
        ext = ext.lower()
        M = Masker(a.words.split(','), a.star)
        try:
            if ext == '.docx':
                dst = base + '_masked.docx'; do_docx(f, dst, M)
            elif ext == '.xlsx':
                dst = base + '_masked.xlsx'; do_xlsx(f, dst, M)
            elif ext == '.pdf':
                dst = base + '_pdf_masked.txt'; do_pdf(f, dst, M)
            elif ext in ('.txt', '.csv', '.md', '.json', '.log', '.tsv'):
                dst = base + '_masked' + ext; do_text(f, dst, M)
            else:
                print(f'건너뜀(지원하지 않는 형식): {f}'); continue
        except Exception as e:
            print(f'실패: {f} - {e}'); continue
        found = ', '.join(f'{k} {v}건' for k, v in M.found.items()) or '찾은 항목 없음'
        print(f'{f} -> {dst}  [{found}]')
        if a.map and M.map:
            json.dump(M.map, open(base + '_map.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        if pw:
            try:
                print('  원본 잠금 ->', lock(f, pw))
            except Exception as e:
                print('  잠금 실패:', e)


if __name__ == '__main__':
    main()
