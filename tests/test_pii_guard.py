"""pii_guard.py 테스트 (가짜 정보만 사용).

실행:  python -m unittest discover tests
필요:  pip install python-docx openpyxl pypdf
"""
import os, re, sys, tempfile, unittest, zipfile
from pathlib import Path

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
import pii_guard as g

# (입력, 기대 결과)
CASES = [
    ('900101-1234567', '[주민등록번호_1]'),
    ('900101-5234567', '[주민등록번호_1]'),            # 외국인등록번호
    ('1234-5678-9012-3456', '[카드번호_1]'),
    ('010-1234-5678', '[휴대전화_1]'),
    ('02-2260-3912', '[전화번호_1]'),
    ('(02)2260-3912', '[전화번호_1]'),
    ('07012345678', '[전화번호_1]'),
    ('+82 10 1234 5678', '[전화번호_1]'),
    ('gildong.hong@example.com', '[이메일_1]'),
    ('123-45-67890', '[사업자번호_1]'),
    ('M12345678', '[여권번호_1]'),
    ('IP는 192.168.0.15.', 'IP는 [IP주소_1].'),
    ('110-123-456789', '[계좌번호_1]'),
    ('2026-09-29', '2026-09-29'),                       # 날짜는 그대로
    ('버전 1.2.3.4.5', '버전 1.2.3.4.5'),                 # IP 아님
    ('연락처(010-1234-5678)', '연락처([휴대전화_1])'),       # 괄호는 그대로 남아야 함
    ('(02-2260-3912)', '([전화번호_1])'),
    ('(010)1234-5678', '[휴대전화_1]'),
    ('+82 (0)2-2260-3912', '[전화번호_1]'),
    ('０１０-１２３４-５６７８', '[휴대전화_1]'),            # 전각 숫자
    ('M123A4567', '[여권번호_1]'),                      # 신형 여권
    ('주문번호 2026102112345', '주문번호 2026102112345'),  # 앞 6자리가 날짜가 아니면 주민번호 아님
    ('서울시 송파구 위례송파로 123, 1706동 804호', '[주소_1]'),
    ('주소: 경기도 성남시 분당구 판교역로 235입니다', '주소: [주소_1]입니다'),
    ('서울 송파구 장지동 123-4번지', '[주소_1]'),               # 지번 주소
    ('송파구 위례송파로 123, 1706동 804호', '[주소_1]'),         # 시·도 없어도 동·호까지 있으면 찾음
    ('송파구 위례송파로 123', '송파구 위례송파로 123'),          # 시·도도 동·호도 없으면 오탐 위험 → 안 찾음
    ('연구 결과로 3배 증가', '연구 결과로 3배 증가'),
    ('경기도 구리시 갈매중앙로 16, 709-702', '[주소_1]'),        # 동-호를 줄여 쓴 형식
    ('구리시 갈매중앙로 16, 709-702', '[주소_1]'),
    ('경기도 갈매중앙로 16, 709-702', '[주소_1]'),               # 시·군·구가 없어도 동·호까지 있으면 찾음
    ('경기 결과로 3배 증가', '경기 결과로 3배 증가'),             # 시·도만 있고 동·호가 없으면 안 찾음
    ('서울시 송파구 위례송파로 123, 2026-10-21 방문', '[주소_1], 2026-10-21 방문'),   # 날짜를 동·호로 보지 않음
]

# 일부만 가리기 (--partial): 업무에 다시 쓸 수 있게 앞뒤만 남김
PARTIAL = [
    ('901023-1234567', '901023-1******'),
    ('010-9249-5149', '010-****-5149'),
    ('서울시 송파구 위례송파로 123, 1706동 804호', '서울시 송파구 위례송파로 ***, ****동 ***호'),
    ('구리시 갈매중앙로 16, 709-702', '구리시 갈매중앙로 **, ***-***'),
    ('1234-5678-9012-3456', '1234-****-****-3456'),
    ('gildong.hong@example.com', 'gi**********@example.com'),
    ('192.168.0.15', '192.168.*.*'),
    ('M12345678', 'M****5678'),
    ('(010)1234-5678', '(010)****-5678'),
]


class Rules(unittest.TestCase):
    def test_cases(self):
        for src, want in CASES:
            with self.subTest(src=src):
                self.assertEqual(g.Masker().text(src), want)

    def test_partial(self):
        for src, want in PARTIAL:
            with self.subTest(src=src):
                self.assertEqual(g.Masker(part=True).text(src), want)
        self.assertEqual(g.Masker(['홍길동'], part=True).text('홍길동'), '홍*동')

    def test_same_value_same_token_and_words(self):
        M = g.Masker(['홍길동'])
        self.assertEqual(M.text('홍길동 010-1234-5678 / 홍길동 010-1234-5678'),
                         '[지정어_1] [휴대전화_1] / [지정어_1] [휴대전화_1]')
        self.assertEqual(M.map['[휴대전화_1]'], '010-1234-5678')

    def test_fullwidth_keeps_original_and_star_hides_it(self):
        M = g.Masker()
        self.assertEqual(M.text('０１０-１２３４-５６７８'), '[휴대전화_1]')
        self.assertEqual(M.map['[휴대전화_1]'], '０１０-１２３４-５６７８')   # 복원하면 원래 글자 그대로
        self.assertEqual(g.Masker(star=True).text('０１０-１２３４-５６７８'), '***-****-****')

    def test_rules_match_web(self):
        """웹(index.html)과 파이썬의 규칙이 똑같은지 확인 → 한쪽만 고치는 실수 방지."""
        page = Path(ROOT, 'index.html').read_text(encoding='utf-8')
        block = page[page.index('const RULES=['):]
        block = block[:block.index('];')]
        web = dict(re.findall(r"\['([^']+)',/(.+?)/g\]", block))
        self.assertEqual(web, {label: rx.pattern for label, rx in g.RULES})


class Files(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.p = lambda name: os.path.join(self.dir.name, name)

    def tearDown(self):
        self.dir.cleanup()

    def test_docx_keeps_format_and_masks_hidden_text(self):
        import docx
        from docx.oxml import parse_xml
        ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run('중요:').bold = True
        p.add_run(' 담당 010-1111-2222 확인 ')
        p.add_run('끝').italic = True
        p = d.add_paragraph('연락처 ')   # 변경 내용 추적으로 지운 번호, 여러 조각으로 나뉜 필드 코드
        p._p.append(parse_xml(f'<w:del {ns} w:id="1" w:author="a"><w:r><w:delText>010-9876-5432</w:delText></w:r></w:del>'))
        for piece in (' HYPERLINK "mailto:kim.yh', '@example.com" '):
            p._p.append(parse_xml(f'<w:r {ns}><w:instrText xml:space="preserve">{piece}</w:instrText></w:r>'))
        # 속성 속 글자: 필드 코드(fldSimple), 변경 추적 작성자 이름
        d.element.body.append(parse_xml(f'<w:p {ns}><w:fldSimple w:instr=" HYPERLINK &quot;mailto:lee.sj@example.com&quot; ">'
                                        f'<w:r><w:t>메일</w:t></w:r></w:fldSimple></w:p>'))
        d.element.body.append(parse_xml(f'<w:p {ns}><w:ins w:id="2" w:author="홍길동"><w:r><w:t>고친 글</w:t></w:r></w:ins></w:p>'))
        d.save(self.p('a.docx'))
        g.do_docx(self.p('a.docx'), self.p('a_masked.docx'), g.Masker(['홍길동']))

        runs = [(r.text, r.bold, r.italic) for r in docx.Document(self.p('a_masked.docx')).paragraphs[0].runs]
        self.assertEqual(runs, [('중요:', True, None), (' 담당 [휴대전화_1] 확인 ', None, None), ('끝', None, True)])
        with zipfile.ZipFile(self.p('a_masked.docx')) as z:
            xml = z.read('word/document.xml').decode()
        for s in ('010-1111-2222', '010-9876-5432', 'kim.yh', 'lee.sj', '홍길동'):
            self.assertNotIn(s, xml)
        self.assertEqual(g.leftover(self.p('a_masked.docx'), ['홍길동']), {})

    def test_docx_people_part(self):
        """Word 2013+가 댓글·변경 추적 작성자를 따로 적는 word/people.xml (이름 + 마이크로소프트 365 계정 이메일)."""
        import docx
        from docx.opc.packuri import PackURI
        from docx.opc.part import Part
        d = docx.Document()
        d.add_paragraph('본문')
        people = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  '<w15:people xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml">'
                  '<w15:person w15:author="홍길동"><w15:presenceInfo w15:providerId="AD" '
                  'w15:userId="S::gildong.hong@example.com::1234abcd"/></w15:person></w15:people>').encode()
        part = Part(PackURI('/word/people.xml'), 'application/vnd.openxmlformats-officedocument.wordprocessingml.people+xml',
                    people, d.part.package)
        d.part.relate_to(part, 'http://schemas.microsoft.com/office/2011/relationships/people')
        d.save(self.p('a.docx'))
        g.do_docx(self.p('a.docx'), self.p('a_masked.docx'), g.Masker(['홍길동']))

        with zipfile.ZipFile(self.p('a_masked.docx')) as z:
            xml = z.read('word/people.xml').decode()
        self.assertNotIn('홍길동', xml)
        self.assertNotIn('gildong.hong@example.com', xml)
        self.assertEqual(g.leftover(self.p('a_masked.docx'), ['홍길동']), {})

    def test_pdf_restricted_opens_but_password_pdf_is_refused(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        for name, user_pw in (('restricted.pdf', ''), ('locked.pdf', 'open-me')):
            w = PdfWriter()
            page = w.add_blank_page(300, 200)
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                                     NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): w._add_object(font)})})
            body = DecodedStreamObject()
            body.set_data(b'BT /F1 12 Tf 10 100 Td (Phone 010-1234-5678) Tj ET')
            page[NameObject('/Contents')] = w._add_object(body)
            w.encrypt(user_password=user_pw, owner_password='owner-secret', permissions_flag=0)   # 복사·인쇄 금지
            with open(self.p(name), 'wb') as f:
                w.write(f)
        g.do_pdf(self.p('restricted.pdf'), self.p('out.txt'), g.Masker())   # 열기 비밀번호 없음 → 처리됨
        self.assertIn('[휴대전화_1]', Path(self.p('out.txt')).read_text(encoding='utf-8'))
        with self.assertRaisesRegex(ValueError, '비밀번호'):
            g.do_pdf(self.p('locked.pdf'), self.p('out2.txt'), g.Masker())

    def test_broken_files_give_korean_messages(self):
        open(self.p('empty.docx'), 'wb').close()
        with open(self.p('old.xlsx'), 'wb') as f:   # 암호화된 Office 파일·옛 형식과 같은 시작 바이트
            f.write(b'\xd0\xcf\x11\xe0' + b'\0' * 100)
        for name, do, word in (('empty.docx', g.do_docx, '손상'), ('old.xlsx', g.do_xlsx, '비밀번호')):
            with self.subTest(name=name), self.assertRaisesRegex(ValueError, word):
                do(self.p(name), self.p('out'), g.Masker())

    def test_xlsx_leading_zero_and_recheck(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 1012345678                      # 서식으로 010-1234-5678처럼 보이는 셀
        ws['A1'].number_format = '000-0000-0000'
        ws['B1'] = 501014123456                    # 0501014123456 (앞자리 0이 지워진 주민번호)
        ws['C1'] = 250000000                       # 금액 2억 5천만 원: 0을 붙이면 02-5000-0000 모양이지만 그대로 둬야 함
        ws['D1'] = 3112345678                      # 금액: 0을 붙이면 031-1234-5678 모양
        ws.title = '010-1234-5678'                 # 시트 이름은 가리지 않음 → 재검사가 알려야 함
        wb.save(self.p('a.xlsx'))
        g.do_xlsx(self.p('a.xlsx'), self.p('a_masked.xlsx'), g.Masker())

        ws = openpyxl.load_workbook(self.p('a_masked.xlsx')).active
        self.assertEqual([c.value for c in ws[1]], ['[휴대전화_1]', '[주민등록번호_1]', 250000000, 3112345678])
        self.assertEqual(g.leftover(self.p('a_masked.xlsx'), []), {'휴대전화': 1})

    def test_text_encodings(self):
        for enc in ('cp949', 'utf-16'):
            with self.subTest(enc=enc):
                with open(self.p('a.csv'), 'wb') as f:
                    f.write('이름,연락처\r\n홍길동,010-1234-5678\r\n'.encode(enc))
                g.do_text(self.p('a.csv'), self.p('a_masked.csv'), g.Masker(['홍길동']))
                raw = Path(self.p('a_masked.csv')).read_bytes()
                self.assertTrue(raw.startswith(b'\xef\xbb\xbf'))   # 엑셀용 BOM
                self.assertEqual(raw[3:].decode('utf-8'), '이름,연락처\r\n[지정어_1],[휴대전화_1]\r\n')


if __name__ == '__main__':
    unittest.main()
